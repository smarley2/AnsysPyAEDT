"""The table a user fills in from a datasheet, one row per part number.

Design: `docs/superpowers/specs/2026-09-04-user-core-catalog-design.md`.

A datasheet family is ten rows off one page, so the reader's job is to import
the nine good rows and say precisely what is wrong with the tenth -- refusing
the whole file for one typo is how a user gives up and goes back to asking for
a rebuild.
"""

from __future__ import annotations

import csv
import io
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from inductor_designer.adapters.catalog.core_table import (
    CORE_TEMPLATE_COLUMNS,
    CoreTableError,
    core_import_template,
    import_core_file,
)
from inductor_designer.domain.catalog_records import CoreFamily, ReviewStatus

# One complete, valid row: a Magnetics Kool Mu toroid's published figures.
_VALID_ROW = {
    "manufacturer": "Magnetics",
    "family": "powder-toroid",
    "partNumber": "0077101A7",
    "materialManufacturer": "Magnetics",
    "materialName": "Kool Mu",
    "materialGrade": "60",
    "coating": "parylene",
    "catalogRevision": "magnetics-powder-2024",
    "sourceUrl": "https://www.mag-inc.com/products/powder-cores/kool-mu-cores",
    "sourcePage": "12",
    "outerDiameterNominalM": "0.0267",
    "outerDiameterMinM": "",
    "outerDiameterMaxM": "",
    "innerDiameterNominalM": "0.0147",
    "innerDiameterMinM": "",
    "innerDiameterMaxM": "",
    "heightNominalM": "0.0112",
    "heightMinM": "",
    "heightMaxM": "",
    "effectiveAreaM2": "6.55e-5",
    "pathLengthM": "0.0635",
    "volumeM3": "4.16e-6",
    "alValueNh": "75.0",
    "reviewStatus": "draft",
    "reviewedBy": "",
}


def _csv_bytes(rows: list[dict[str, str]]) -> bytes:
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=list(CORE_TEMPLATE_COLUMNS))
    writer.writeheader()
    writer.writerows(rows)
    return buffer.getvalue().encode("utf-8")


def test_the_template_columns_are_the_schema_fields() -> None:
    """The template and the validator must never disagree about what a column
    means, so the columns ARE the schema's field names -- with the material
    object and each dimension flattened, since a spreadsheet has no nesting."""
    assert CORE_TEMPLATE_COLUMNS[:3] == ("manufacturer", "family", "partNumber")
    for expected in (
        "materialManufacturer",
        "sourceUrl",
        "sourcePage",
        "outerDiameterNominalM",
        "innerDiameterMinM",
        "heightMaxM",
        "effectiveAreaM2",
        "pathLengthM",
        "volumeM3",
        "alValueNh",
        "reviewStatus",
    ):
        assert expected in CORE_TEMPLATE_COLUMNS


@pytest.mark.parametrize("file_format", ["csv", "xlsx"])
def test_the_template_downloads_empty_with_a_header(file_format: str) -> None:
    download = core_import_template(file_format)
    assert download.filename.endswith(file_format)
    assert download.data

    result = import_core_file(download.filename, download.data)
    # A template with no rows imports nothing and complains about nothing:
    # it is a starting point, not a mistake.
    assert result.records == ()
    assert result.rejections == ()


def test_an_unknown_template_format_is_refused() -> None:
    with pytest.raises(ValueError, match="csv"):
        core_import_template("pdf")


def test_a_filled_row_becomes_the_record_that_was_typed() -> None:
    result = import_core_file("cores.csv", _csv_bytes([_VALID_ROW]))

    assert result.rejections == ()
    (core,) = result.records
    assert core.part_number == "0077101A7"
    assert core.family is CoreFamily.POWDER_TOROID
    assert core.material.name == "Kool Mu"
    assert core.source_page == 12
    assert core.outer_diameter.nominal_m == pytest.approx(0.0267)
    # An empty tolerance cell is "not published", not zero.
    assert core.outer_diameter.min_m is None
    assert core.al_value_nh == pytest.approx(75.0)
    assert core.review_status is ReviewStatus.DRAFT


def test_nine_good_rows_survive_one_bad_row() -> None:
    rows = []
    for index in range(10):
        row = dict(_VALID_ROW)
        row["partNumber"] = f"CORE-{index}"
        rows.append(row)
    rows[4]["alValueNh"] = "not a number"

    result = import_core_file("family.csv", _csv_bytes(rows))

    assert [core.part_number for core in result.records] == [
        f"CORE-{index}" for index in range(10) if index != 4
    ]
    (rejection,) = result.rejections
    # Row 6: the header is row 1, so the fifth data row is the sixth line --
    # the number the user sees in their spreadsheet.
    assert rejection.row == 6
    assert "alValueNh" in rejection.reason


def test_a_value_the_schema_refuses_is_rejected_with_its_reason() -> None:
    """The catalog schema is the authority on what a core may claim, so the
    reader validates against it rather than re-deciding the rules."""
    row = dict(_VALID_ROW)
    row["family"] = "e-core"

    result = import_core_file("cores.csv", _csv_bytes([row]))

    assert result.records == ()
    (rejection,) = result.rejections
    assert rejection.row == 2
    assert "family" in rejection.reason


def test_a_missing_column_refuses_the_whole_file() -> None:
    """Unlike a bad row, a wrong header means every row is suspect -- there is
    no way to tell which value was meant for which field."""
    rows = [dict(_VALID_ROW)]
    buffer = io.StringIO()
    columns = [name for name in CORE_TEMPLATE_COLUMNS if name != "pathLengthM"]
    writer = csv.DictWriter(buffer, fieldnames=columns, extrasaction="ignore")
    writer.writeheader()
    writer.writerows(rows)

    with pytest.raises(CoreTableError, match="pathLengthM"):
        import_core_file("cores.csv", buffer.getvalue().encode("utf-8"))


def test_a_formula_cell_is_refused_rather_than_evaluated() -> None:
    """An imported A_L has to be a number a person read off a datasheet. A
    formula is a value this application cannot see, cannot cite, and would
    silently re-evaluate differently in another spreadsheet."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(list(CORE_TEMPLATE_COLUMNS))
    sheet.append([_VALID_ROW[name] for name in CORE_TEMPLATE_COLUMNS])
    sheet.cell(row=2, column=list(CORE_TEMPLATE_COLUMNS).index("alValueNh") + 1).value = (
        "=SUM(1,2)"
    )
    buffer = io.BytesIO()
    workbook.save(buffer)

    with pytest.raises(CoreTableError, match="formula"):
        import_core_file("cores.xlsx", buffer.getvalue())


def test_an_xlsx_round_trip_reads_the_same_record_as_the_csv() -> None:
    """The two formats are one reader's two front doors; a value must not mean
    different things depending on which one the user chose."""
    download = core_import_template("xlsx")
    workbook = load_workbook(io.BytesIO(download.data))
    sheet = workbook.active
    sheet.append([_VALID_ROW[name] for name in CORE_TEMPLATE_COLUMNS])
    buffer = io.BytesIO()
    workbook.save(buffer)

    from_xlsx = import_core_file("cores.xlsx", buffer.getvalue()).records
    from_csv = import_core_file("cores.csv", _csv_bytes([_VALID_ROW])).records
    assert from_xlsx == from_csv


def test_a_file_that_is_not_a_table_at_all_is_refused(tmp_path: Path) -> None:
    with pytest.raises(CoreTableError):
        import_core_file("cores.xlsx", b"not a workbook")
